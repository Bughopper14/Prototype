import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
      <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideV w:val="none"/>
      <w:left w:val="none"/>
      <w:right w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_callout(doc, text, title=None, border_color="002060", fill_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
      <w:left w:val="single" w:sz="36" w:color="{border_color}"/>
      <w:top w:val="none"/>
      <w:right w:val="none"/>
      <w:bottom w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(title + "\n")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(0, 32, 96)
    
    run_text = p.add_run(text)
    run_text.italic = True
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = RGBColor(50, 60, 75)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_word_document():
    doc = docx.Document()
    
    # Page setup (A4 standard)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Configure Header & Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("PT LIUGONG FINANCE INDONESIA  |  ORIGINATION WORKSPACE & CRM (LOS-CRM)")
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(120, 130, 140)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Dokumen Spesifikasi Bisnis & Panduan Eksekutif Direksi  •  Rahasia & Hak Milik Perusahaan")
    frun.font.name = "Calibri"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(120, 130, 140)

    # -------------------------------------------------------------
    # COVER PAGE
    # -------------------------------------------------------------
    logo_path = r"D:\CRM App2\CRM App2\public\excel_assets\image2.jpg"
    if os.path.exists(logo_path):
        logo_p = doc.add_paragraph()
        logo_p.paragraph_format.space_before = Pt(20)
        logo_p.paragraph_format.space_after = Pt(24)
        logo_run = logo_p.add_run()
        logo_run.add_picture(logo_path, width=Inches(2.6))
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(6)
    run_t = p_title.add_run("SISTEM ORIGINASI PEMBIAYAAN & MANAJEMEN NASABAH (LOS-CRM)")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(24)
    run_t.bold = True
    run_t.font.color.rgb = RGBColor(0, 32, 96) # LiuGong Navy

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("DOKUMEN SPESIFIKASI BISNIS, STRATEGI PIPELINE 4-TAHAP & PANDUAN EKSEKUTIF DIREKSI")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.bold = True
    run_sub.font.color.rgb = RGBColor(237, 125, 49) # LiuGong Orange

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    add_callout(doc, 
        "\"Platform LOS-CRM ini dirancang sebagai mesin akselerasi bisnis (Growth Engine) yang mentransformasikan proses pembiayaan alat berat tradisional menjadi Fintech-Grade Enterprise Platform. Memangkas Turnaround Time (TAT) dari 14 hari menjadi 48 jam, mengunci mitigasi risiko kredit sejak detik pertama, dan mengamankan kepatuhan penuh terhadap regulasi OJK & UU Pelindungan Data Pribadi.\"",
        title="EXECUTIVE SUMMARY UNTUK PRESIDEN DIREKTUR & BOARD OF DIRECTORS")

    # Metadata table on cover
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="D3D3D3")
    
    col_w = [Inches(2.2), Inches(4.3)]
    meta_rows = [
        ("Nama Proyek & Sistem", "LiuGong Finance Origination Workspace & CRM (LOS-CRM)"),
        ("Entitas Pembiayaan", "PT LiuGong Finance Indonesia (Captive Finance)"),
        ("Disusun Oleh", "Latom (Lead System Owner & Technical In-Charge)"),
        ("Target Presentasi", "Bapak Presiden Direktur & Jajaran Komite Kredit"),
        ("Status Pengembangan", "Fase 1-2 Selesai (100% Tested) — Siap Integrasi Pefindo & Multi-OCR"),
        ("Tanggal Efektif Dokumen", "19 Oktober 2026  |  Versi 3.2.0 (Executive Master)")
    ]
    for i, (k, v) in enumerate(meta_rows):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_w[0], col_w[1]
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0, 32, 96)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(40, 50, 60)
        if i % 2 == 1:
            set_cell_background(c0, "F8FAFC")
            set_cell_background(c1, "F8FAFC")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 1: VISI STRATEGIS & LATAR BELAKANG
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("1. LATAR BELAKANG, VISI BISNIS & SASARAN DIREKSI")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Di tengah ekspansi sektor pertambangan nikel, batubara, perkebunan kelapa sawit, dan pembangunan infrastruktur "
        "di Indonesia, kecepatan penyediaan alat berat merupakan faktor penentu kemenangan bisnis. Pelanggan kontraktor "
        "dan pemilik tambang menuntut kepastian persetujuan pembiayaan dalam hitungan hari. "
        "Sistem lama yang bergantung pada lembar Excel lepas dan sirkulasi berkas fisik menimbulkan hambatan berupa proses "
        "persetujuan lambat (10–14 hari kerja), risiko kesalahan hitung angsuran, serta keterlambatan deteksi nasabah berisiko tinggi."
    )

    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(10)
    p2.add_run(
        "Platform LOS-CRM ini dibangun untuk menjawab tiga sasaran strategis Presiden Direktur:\n"
        "• Kecepatan Layanan (Speed of Delivery): Memangkas total siklus persetujuan kredit (TAT) menjadi maksimal 48 jam.\n"
        "• Presisi Mitigasi Risiko (Risk Precision from Second Zero): Memfilter debitur bermasalah sejak detik pertama melalui otomatisasi Pefindo Biro Kredit dan OCR identitas resmi.\n"
        "• Efisiensi & Kepatuhan Penuh (Zero-Error & Compliance): Mengeliminasi human error pada kalkulasi anuitas, validasi seimbang neraca 100%, serta audit trail yang memenuhi regulasi OJK & UU PDP No. 27/2022."
    )

    # -------------------------------------------------------------
    # SECTION 2: ALUR BESAR STRATEGIS 4-TAHAP
    # -------------------------------------------------------------
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run("2. PIPA STRATEGIS 4-TAHAP (THE 4-STAGE STRATEGIC PIPELINE)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Alur kerja pembiayaan pada aplikasi dirancang dalam 4 tahapan progresif yang linear, transparan, dan akuntabel:\n"
        "FAP STAGE  ⟶  [PEFINDO OCR GATEWAY]  ⟶  PRE-ANALYSIS  ⟶  ICLF  ⟶  PROPOSAL"
    )

    flow_table = doc.add_table(rows=6, cols=3)
    flow_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow_table.autofit = False
    set_table_borders(flow_table, color="D3D3D3")
    
    col_w_flow = [Inches(1.8), Inches(2.2), Inches(2.5)]
    headers = ["Tahapan Sistem", "Fitur Unggulan & Teknologi", "Nilai Strategis bagi Manajemen"]
    for j, htext in enumerate(headers):
        cell = flow_table.rows[0].cells[j]
        cell.width = col_w_flow[j]
        set_cell_background(cell, "002060")
        set_cell_margins(cell, 80, 80, 100, 100)
        p_hdr = cell.paragraphs[0]
        r_hdr = p_hdr.add_run(htext)
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(10)
        r_hdr.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)

    flow_rows = [
        ("1. FAP STAGE\n(Intake & Profiling)", 
         "• Multi-Document AI/OCR Suite\n  (NIB, e-KTP, Paspor, NPWP, E-KITAS)\n• Form 5-Modul Accordion\n• Digital Twin Excel FAP 118-Baris",
         "Onboarding nasabah super cepat. Mengurangi 80% waktu pengetikan dan meniadakan salah input nomor identitas."),
        ("GATEWAY INTERKONEKSI\n(Pefindo Biro Kredit)", 
         "• Ingestion Data ke Pefindo OCR/API\n• Penarikan SLIK OJK 24 Bulan\n• Credit Score (250–900)\n• Total Eksposur Utang Perbankan",
         "Mengetahui rekam jejak kolektibilitas (Kol 1–5) secara seketika sebelum melangkah ke analisis mendalam."),
        ("2. PRE-ANALYSIS\n(Risk Gatekeeping)", 
         "• Scoring Kelayakan Awal\n• Pemetaan Sektor Industri (Mining, dll)\n• Triase Go / Review / No-Go",
         "Early warning filter. Menghemat sumber daya analis agar tidak memproses debitur berisiko macet."),
        ("3. ICLF\n(Limit Setting)", 
         "• Batas Plafon Tunggal & Grup Usaha\n• Evaluasi BMPP Regulasi OJK\n• Struktur Rekomendasi Uang Muka (DP)",
         "Pencegahan konsentrasi kredit berlebih dan memastikan kecukupan permodalan debitur."),
        ("4. PROPOSAL\n(Dossier & Putusan)", 
         "• Inventaris Unit Alat Berat LiuGong\n• Anuitas Bunga Efektif Menurun\n• Neraca 3 Tahun 100% Seimbang\n• Putusan Komite Kredit Digital",
         "Dossier lengkap siap sidang. Seluruh putusan kredit terkunci aman dengan jejak audit abadi.")
    ]

    for i, (col1, col2, col3) in enumerate(flow_rows):
        row = flow_table.rows[i+1]
        for c_idx, text in enumerate([col1, col2, col3]):
            c = row.cells[c_idx]
            c.width = col_w_flow[c_idx]
            set_cell_margins(c, 70, 70, 90, 90)
            p_cell = c.paragraphs[0]
            p_cell.paragraph_format.line_spacing = 1.15
            p_cell.paragraph_format.space_after = Pt(2)
            r_c = p_cell.add_run(text)
            r_c.font.name = "Calibri"
            r_c.font.size = Pt(9)
            if c_idx == 0:
                r_c.bold = True
                r_c.font.color.rgb = RGBColor(0, 32, 96)
            else:
                r_c.font.color.rgb = RGBColor(40, 50, 60)
            if i % 2 == 1:
                set_cell_background(c, "F8FAFC")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 3: MULTI-DOC OCR & PEFINDO INTEGRATION
    # -------------------------------------------------------------
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    r = h3.add_run("3. DEEP DIVE: MULTI-DOCUMENT AI/OCR & PEFINDO GATEWAY")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    sub31 = doc.add_paragraph()
    sub31.paragraph_format.space_before = Pt(8)
    sub31.paragraph_format.space_after = Pt(4)
    r = sub31.add_run("3.1. Suite AI/OCR Multidokumen pada FAP Stage")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Pada tahap intake FAP, aplikasi dilengkapi kecerdasan buatan OCR (Optical Character Recognition) mutakhir "
        "yang berjalan langsung di sisi peramban pengguna (Client-Side WASM Engine). Hal ini menjamin privasi data nasabah "
        "karena berkas identitas tidak pernah dikirim ke pihak ketiga secara mentah. Dokumen yang didukung meliputi:\n"
        "1. NIB (Nomor Induk Berusaha): Mengekstrak 13 digit NIB, nama perusahaan, alamat kantor, kode pos, dan tanggal penetapan dari sertifikat resmi OSS RBA Kemeninvestasi/BKPM.\n"
        "2. e-KTP (Kartu Tanda Penduduk WNI): Mengekstrak NIK 16 digit, nama lengkap, tempat/tanggal lahir, jenis kelamin, dan alamat tempat tinggal Direksi atau Pemegang Saham WNI.\n"
        "3. PASSPORT & E-KITAS (Expatriate / Asing): Mengekstrak nomor paspor, kewarganegaraan, masa berlaku, dan dokumen izin tinggal bagi pemodal atau manajemen berkewarganegaraan asing.\n"
        "4. NPWP (Nomor Pokok Wajib Pajak): Validasi instan nomor pokok pajak badan usaha maupun perorangan (format 15 digit lama maupun 16 digit NIK-NPWP)."
    )

    sub32 = doc.add_paragraph()
    sub32.paragraph_format.space_before = Pt(10)
    sub32.paragraph_format.space_after = Pt(4)
    r = sub32.add_run("3.2. Jembatan Cerdas Antara FAP Stage dan Pre-Analysis: Pefindo Biro Kredit")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Begitu identitas badan usaha dan para pengurus kunci tervalidasi pada FAP Stage, sistem secara otomatis "
        "dapat menembakkan data tersebut ke Pefindo Biro Kredit (IdScore / SLIK OJK Ingestion Bridge). "
        "Langkah ini menghadirkan profil intelijen finansial yang komprehensif dalam hitungan detik:"
    )

    add_callout(doc,
        "MANFAAT STRATEGIS PEFINDO BIRO KREDIT:\n"
        "• Skor Kredit Obyektif (Pefindo Score 250–900): Mengetahui tingkat risiko gagal bayar (Probability of Default) nasabah secara presisi.\n"
        "• Riwayat Kolektibilitas 24 Bulan: Mengetahui histori cicilan debitur di bank atau leasing lain—apakah selalu Lancar (Kol 1) atau pernah Nunggak/Macet (Kol 2–5).\n"
        "• Total Fasilitas Aktif se-Indonesia: Menghitung total beban angsuran bulanan yang sedang ditanggung debitur saat ini.\n"
        "• Deteksi Dini Restrukturisasi & Gugatan: Menghindarkan LiuGong Finance dari nasabah yang sedang dalam sengketa hukum perdata atau pailit.",
        title="THE PEFINDO INTELLIGENCE ADVANTAGE")

    # -------------------------------------------------------------
    # SECTION 4: USER INTERFACE & WORKSPACE PARITY
    # -------------------------------------------------------------
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r = h4.add_run("4. KEMUDAHAN OPERASIONAL: WORKSPACE SPLIT 50-50 & DIGITAL TWIN EXCEL")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Salah satu keunggulan terbesar dari aplikasi ini adalah desain antarmuka Split 50-50 yang menyandingkan kemudahan "
        "pengisian formulir modern di sisi kiri dengan pratinjau lembar resmi Excel FAP 118-baris di sisi kanan:\n"
        "• Form 5-Modul Accordion: Memecah lebih dari 35 field input menjadi 5 modul bernomor rapi (01. Identitas, 02. Domisili, 03. Profil Usaha & NIB OCR, 04. PIC, 05. Proyek Tambang/Konstruksi).\n"
        "• Bidirectional Auto-Sync: Saat tim marketing mengetik di kiri, sel di dokumen Excel kanan langsung terisi secara live dengan efek highlight biru yang berkedip halus. Sebaliknya, mengklik chip navigasi di lembar dokumen kanan otomatis membuka modul form di kiri.\n"
        "• Preservasi Format Dokumen Resmi: Manajemen dan pimpinan cabang tetap melihat format dokumen FAP yang sudah mereka kenal puluhan tahun tanpa merasa asing."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 5: GOVERNANCE & VALIDATION INTEGRITY
    # -------------------------------------------------------------
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(6)
    r = h5.add_run("5. TATA KELOLA KREDIT, INTEGRITAS FORMULA & KEPATUHAN OJK")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Aplikasi menerapkan mekanisme validasi ketat dan berlapis untuk mengamankan portofolio kredit perusahaan:"
    )

    gov_points = [
        ("Validasi Kepemilikan Saham 100% Persis", 
         "Sistem secara deterministik menolak pengajuan permohonan kredit jika total porsi kepemilikan saham seluruh pemegang saham tidak berjumlah persis 100.00%. Hal ini menutup celah kepemilikan fiktif dan mematuhi prinsip anti-pencucian uang (APU-PPT OJK)."),
        ("Integritas Keseimbangan Neraca (Balance Sheet Equilibrium)", 
         "Menggunakan kalkulasi sen integer (BigInt) presisi tinggi untuk memeriksa 3 tahun laporan keuangan: Total Aset wajib sama persis dengan Total Liabilitas + Ekuitas. Laporan yang tidak seimbang langsung ditolak oleh sistem."),
        ("Kalkulasi Anuitas Bunga Efektif Otomatis", 
         "Menghitung angsuran bulanan menggunakan rumus reducing balance annuity aktuaria resmi, baik untuk skema In Arrear maupun In Advance, meniadakan selisih hitung manual antara tim marketing dan tim finance."),
        ("Tata Kelola Akta Legalitas & Notaris", 
         "Pemisahan akta aktif dan arsip historis, dilengkapi editor teks TinyMCE dengan penomoran klausul kontinu dan matriks perbandingan perubahan modal/pengurus secara berdampingan."),
        ("Jejak Audit Abadi (Immutable Append-Only Audit Trail)", 
         "Tabel audit logs dilindungi oleh trigger basis data tingkat kernel: tidak ada seorang pun (termasuk administrator IT) yang dapat mengubah atau menghapus riwayat audit. Menjamin kesiapan 100% saat diaudit oleh regulator OJK."),
        ("Kepatuhan UU PDP No. 27/2022", 
         "Enkripsi dokumen agunan dan identitas di AWS S3 privat menggunakan tanda tangan digital dengan kedaluwarsa 15 menit, mencegah kebocoran data rahasia debitur.")
    ]

    for title, desc in gov_points:
        p_pt = doc.add_paragraph()
        p_pt.paragraph_format.line_spacing = 1.15
        p_pt.paragraph_format.space_after = Pt(4)
        run_pt_t = p_pt.add_run(f"• {title}: ")
        run_pt_t.bold = True
        run_pt_t.font.name = "Calibri"
        run_pt_t.font.size = Pt(10)
        run_pt_t.font.color.rgb = RGBColor(0, 32, 96)
        run_pt_d = p_pt.add_run(desc)
        run_pt_d.font.name = "Calibri"
        run_pt_d.font.size = Pt(10)
        run_pt_d.font.color.rgb = RGBColor(40, 50, 60)

    # -------------------------------------------------------------
    # SECTION 6: METRIK KEBERHASILAN & ROADMAP
    # -------------------------------------------------------------
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(16)
    h6.paragraph_format.space_after = Pt(6)
    r = h6.add_run("6. METRIK KEBERHASILAN BISNIS, KPI & ROADMAP IMPLEMENTASI")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    kpi_table = doc.add_table(rows=6, cols=3)
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    kpi_table.autofit = False
    set_table_borders(kpi_table, color="D3D3D3")
    
    col_w_kpi = [Inches(2.5), Inches(1.8), Inches(2.2)]
    kpi_headers = ["Indikator Kinerja Utama (KPI)", "Sebelum Sistem Baru", "Target Sistem LOS-CRM"]
    for j, htext in enumerate(kpi_headers):
        cell = kpi_table.rows[0].cells[j]
        cell.width = col_w_kpi[j]
        set_cell_background(cell, "002060")
        set_cell_margins(cell, 70, 70, 90, 90)
        p_hdr = cell.paragraphs[0]
        r_hdr = p_hdr.add_run(htext)
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(9.5)
        r_hdr.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)

    kpi_rows = [
        ("Turnaround Time (TAT) Siklus Kredit", "10 – 14 Hari Kerja", "Maksimal 48 Jam (2 Hari)"),
        ("Waktu Onboarding & Input Data Nasabah", "4 – 6 Jam", "≤ 30 Menit (Multi-Doc OCR)"),
        ("Inquiry Riwayat Kredit SLIK / Pefindo", "1 – 2 Hari Kerja", "Real-Time / Hitungan Detik"),
        ("Tingkat Retur Berkas (Defect Rate)", "35% – 40% aplikasi", "≤ 3% permohonan"),
        ("Akurasi Laporan Keuangan & Saham", "Rentan salah hitung", "100% Zero-Error Terverifikasi")
    ]

    for i, (kpi, old, new) in enumerate(kpi_rows):
        row = kpi_table.rows[i+1]
        for c_idx, text in enumerate([kpi, old, new]):
            c = row.cells[c_idx]
            c.width = col_w_kpi[c_idx]
            set_cell_margins(c, 60, 60, 80, 80)
            p_cell = c.paragraphs[0]
            r_c = p_cell.add_run(text)
            r_c.font.name = "Calibri"
            r_c.font.size = Pt(9)
            if c_idx == 0:
                r_c.bold = True
                r_c.font.color.rgb = RGBColor(0, 32, 96)
            elif c_idx == 2:
                r_c.bold = True
                r_c.font.color.rgb = RGBColor(237, 125, 49)
            if i % 2 == 1:
                set_cell_background(c, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Formal Signatures
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    set_table_borders(sig_table, color="FFFFFF")
    
    sig_col_w = [Inches(3.2), Inches(3.3)]
    row0 = sig_table.rows[0]
    p_s1 = row0.cells[0].paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s1 = p_s1.add_run("Disiapkan & Diajukan Oleh:\nLead System Owner & Technical In-Charge\n\n\n\n\n")
    r_s1.font.name = "Calibri"
    r_s1.font.size = Pt(9.5)
    r_s1_b = p_s1.add_run("LATOM\n")
    r_s1_b.bold = True
    r_s1_b.font.name = "Calibri"
    r_s1_b.font.size = Pt(10.5)
    r_s1_b.font.color.rgb = RGBColor(0, 32, 96)
    p_s1.add_run("PT LiuGong Finance Indonesia")

    p_s2 = row0.cells[1].paragraphs[0]
    p_s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s2 = p_s2.add_run("Ditelaah & Disetujui Oleh:\nPresiden Direktur\n\n\n\n\n")
    r_s2.font.name = "Calibri"
    r_s2.font.size = Pt(9.5)
    r_s2_b = p_s2.add_run("PRESIDEN DIREKTUR\n")
    r_s2_b.bold = True
    r_s2_b.font.name = "Calibri"
    r_s2_b.font.size = Pt(10.5)
    r_s2_b.font.color.rgb = RGBColor(0, 32, 96)
    p_s2.add_run("PT LiuGong Finance Indonesia")

    output_path = r"D:\CRM App2\CRM App2\LIUGONG_CRM_BUSINESS_SPECIFICATION_PRESDIR.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path} (size: {os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    build_word_document()
