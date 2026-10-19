import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
      <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideV w:val="none"/>
      <w:left w:val="none"/>
      <w:right w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_callout(doc, text, title=None, border_color="002060", fill_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
      <w:left w:val="single" w:sz="36" w:color="{border_color}"/>
      <w:top w:val="none"/>
      <w:right w:val="none"/>
      <w:bottom w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(title + "\n")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(0, 32, 96)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(40, 50, 65)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_figure(doc, img_path, caption, width=Inches(6.2)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run = p_img.add_run()
        run.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(f"Gambar: {caption}")
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(100, 116, 139)

def build_word_document():
    doc = docx.Document()
    
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # -------------------------------------------------------------
    # HEADER FORMAL: LOGO & JUDUL
    # -------------------------------------------------------------
    logo_path = r"D:\CRM App2\CRM App2\public\excel_assets\image2.jpg"
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_logo.paragraph_format.space_after = Pt(4)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.5))

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("PROPOSAL SPESIFIKASI BISNIS & ROADMAP SISTEM")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0, 32, 96)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("ORIGINATION WORKSPACE & CRM PLATFORM (LOS-CRM) · PT LIUGONG FINANCE INDONESIA")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(237, 125, 49)

    add_callout(
        doc,
        "DOKUMEN PRESENTASI EKSEKUTIF UNTUK PRESIDEN DIREKTUR:\n"
        "Dokumen ini membedah Roadmap Strategis Pengembangan Sistem 2026–2027 dengan fokus utama pada rilis Q4 2026 (FAP STAGE: Customer & PIC hingga Stakeholders + Tab Dokumen). "
        "Dilengkapi rincian operasional langkah demi langkah (SOP 6 Tahap) penginputan data oleh tim Marketing, integrasi OCR NIB OSS RBA, Digital Twin Excel 118-baris, validasi saham 100%, serta gambaran umum fase lanjutan 2027.",
        title="RINGKASAN STRATEGIS DIREKSI"
    )

    # -------------------------------------------------------------
    # TABEL METADATA
    # -------------------------------------------------------------
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="B0C4DE")
    col_w = [Inches(2.2), Inches(4.3)]
    
    meta_data = [
        ("Nama Inisiatif", "LiuGong Finance Origination & Customer Relationship Management (LOS-CRM)"),
        ("Sponsor Proyek", "Presiden Direktur · PT LiuGong Finance Indonesia"),
        ("Lead System Owner", "Latom (CTO & Head of Digital Transformation)"),
        ("Target Rilis Q4 2026", "FAP Stage (Customer & PIC, Legal, Stakeholders, Tab Documents 13 Checklist)"),
        ("Horizon Roadmap 2027", "Q1: Pre-Analysis & Pefindo OCR | Q2: ICLF & CDE Financial | Q3: CDE Bank Statement | Q4: Proposal Final"),
        ("Status Kompilasi & QA", "Production Ready · 27/27 Test Suites PASS · 0 Console Error")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_w[0], col_w[1]
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_background(c0, "F0F4F8")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0, 32, 96)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(40, 50, 60)

    doc.add_page_break()

    # -------------------------------------------------------------
    # BAB 1: ROADMAP STRATEGIS 2026–2027 (PHASED LIFECYCLE)
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("BAB 1: ROADMAP STRATEGIS PENGEMBANGAN SISTEM (2026 – 2027)")
    r.font.name = "Calibri"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Untuk memastikan kehati-hatian perbankan (prudential financing) dan kesiapan adopsi di lapangan, "
        "transformasi sistem LOS-CRM LiuGong dibagi ke dalam 5 kuartal strategis yang terukur:"
    )

    # Tabel Roadmap
    road_table = doc.add_table(rows=6, cols=3)
    road_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    road_table.autofit = False
    set_table_borders(road_table, color="D3D3D3")
    col_w_road = [Inches(1.4), Inches(2.5), Inches(2.6)]
    headers_road = ["Periode Rilis", "Modul & Fitur Kunci", "Sasaran Strategis & Nilai Bisnis"]
    for j, htext in enumerate(headers_road):
        cell = road_table.rows[0].cells[j]
        cell.width = col_w_road[j]
        set_cell_background(cell, "002060")
        set_cell_margins(cell, 60, 60, 80, 80)
        p_hdr = cell.paragraphs[0]
        r_hdr = p_hdr.add_run(htext)
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(9.5)
        r_hdr.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)

    roadmap_rows = [
        ("Q4 2026\n(FOKUS UTAMA)",
         "• FAP STAGE RELEASE:\n  - Customer & PIC (5 Modul Accordion)\n  - Legal Deeds (Akta & SK Menkumham)\n  - Stakeholders (Susunan Pengurus & Saham)\n  - Documents Tab (13 Checklist Berkas)\n  - Live Excel Mirror 118-Baris & OCR NIB",
         "Standarisasi onboarding nasabah, peniadaan typo data legalitas via OCR NIB, validasi kepemilikan modal mutlak 100%, serta sentralisasi berkas digital."),
        
        ("Q1 2027",
         "• PRE-ANALYSIS + INTEGRASI PEFINDO OCR:\n  - Early Risk Screening (Go / Review / No-Go)\n  - Pefindo Credit Score & Ingestion Bridge\n  - OCR NPWP & SPT Tahunan Langsung Mengisi Financial Statements",
         "Penyaringan dini debitur berisiko tinggi. Ekstraksi otomatis laporan keuangan dari dokumen pajak tanpa pengetikan manual oleh analis."),
        
        ("Q2 2027",
         "• ICLF + CDE (FINANCIAL):\n  - Initial Credit Limit Facility (BMPP OJK)\n  - Credit Decision Engine (CDE) Phase 1\n  - Analisis Rasio Likuiditas & Solvabilitas Otomatis",
         "Penetapan pagu batas atas kredit debitur tunggal/grup usaha dan kalkulasi kesehatan neraca secara komputasi otomatis."),
        
        ("Q3 2027",
         "• CDE (BANK STATEMENT) PHASE 2:\n  - Parser Mutasi Rekening Koran Bank Otomatis\n  - Analisis Saldo Mengendap & Cashflow Riil\n  - Deteksi Cek Kosong & Transaksi Anomali",
         "Mengetahui kapasitas angsuran riil debitur dari rekam jejak mutasi bank operasional 3–6 bulan secara objektif."),
        
        ("Q4 2027",
         "• PROPOSAL PEMBIAYAAN FINAL:\n  - Inventaris Unit Alat Berat LiuGong\n  - Anuitas Bunga Efektif Menurun Resmi\n  - Dossier Komprehensif Siap Sidang\n  - Persetujuan Komite Kredit Digital",
         "Siklus kredit penuh selesai (end-to-end). Penerbitan SPPK otomatis dan pencairan unit alat berat LiuGong ke pelanggan.")
    ]

    for i, (col1, col2, col3) in enumerate(roadmap_rows):
        row = road_table.rows[i+1]
        for c_idx, text in enumerate([col1, col2, col3]):
            c = row.cells[c_idx]
            c.width = col_w_road[c_idx]
            set_cell_margins(c, 50, 50, 80, 80)
            p_cell = c.paragraphs[0]
            p_cell.paragraph_format.line_spacing = 1.15
            p_cell.paragraph_format.space_after = Pt(2)
            r_c = p_cell.add_run(text)
            r_c.font.name = "Calibri"
            r_c.font.size = Pt(8.5)
            if c_idx == 0:
                r_c.bold = True
                r_c.font.color.rgb = RGBColor(0, 32, 96)
            else:
                r_c.font.color.rgb = RGBColor(40, 50, 60)
            if i == 0:
                set_cell_background(c, "EFF6FF")
            elif i % 2 == 1:
                set_cell_background(c, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Embed Diagram Pipa 4-Tahap
    add_image_figure(doc, r"D:\CRM App2\CRM App2\public\diagram_pipa_strategis_4_tahap.png",
                     "Diagram Arsitektur Pipa Pembiayaan 4-Tahap & Interkoneksi Pefindo OCR")

    doc.add_page_break()

    # -------------------------------------------------------------
    # BAB 2: PANDUAN OPERASIONAL STEP-BY-STEP FAP STAGE UNTUK MARKETING (SOP 6 LANGKAH)
    # -------------------------------------------------------------
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    r = h2.add_run("BAB 2: PANDUAN OPERASIONAL STEP-BY-STEP FAP STAGE (SOP 6 LANGKAH MARKETING)")
    r.font.name = "Calibri"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Pada kuartal Q4 2026, sistem diresmikan dengan fokus utama pada FAP STAGE. "
        "Berikut adalah Standard Operating Procedure (SOP) langkah demi langkah yang dijalankan oleh Marketing Officer (MKT) "
        "di seluruh kantor cabang LiuGong saat menginput permohonan pembiayaan baru:"
    )

    add_image_figure(doc, r"D:\CRM App2\screenshots_presdir\fap_stage_tabs_with_document.png",
                     "Tampilan Workspace FAP Stage Aktif Menampilkan 4 Tab Terstruktur Lengkap dengan Tab Dokumen (Live System)")

    # LANGKAH 1
    sub_l1 = doc.add_paragraph()
    sub_l1.paragraph_format.space_before = Pt(8)
    sub_l1.paragraph_format.space_after = Pt(2)
    r = sub_l1.add_run("Langkah 1: Inisiasi Permohonan Baru & Pendaftaran Registrasi")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "1. Login Pengguna: Marketing masuk ke sistem menggunakan akun resmi cabang (misal: marketing@liugong.local).\n"
        "2. Buka Origination Workspace: Memilih menu 'Applications' pada sidebar kiri.\n"
        "3. Inisiasi Aplikasi: Klik tombol '+ New application' (atau membuka draft berjalan). "
        "Sistem secara otomatis mengalokasikan nomor register resmi sesuai format standar: [Nomor Urut]/AF/[Bulan Romawi]/[Tahun] (contoh: 003/AF/IX/2026) "
        "dan menetapkan tahap awal pada 'FAP Stage' berstatus 'Draft'."
    )

    # LANGKAH 2
    sub_l2 = doc.add_paragraph()
    sub_l2.paragraph_format.space_before = Pt(8)
    sub_l2.paragraph_format.space_after = Pt(2)
    r = sub_l2.add_run("Langkah 2: Pengisian Tab 1 (Customer & PIC) — 5 Modul Accordion & Digital Twin Excel")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "Marketing melengkapi data identitas calon debitur pada formulir sebelah kiri yang terbagi ke dalam 5 modul kartu accordion bernomor:\n"
        "• Modul 01 - Identitas & Klasifikasi Debitur:\n"
        "  - Memilih Kategori Pelanggan: 'Customer Baru' (New Customer) atau 'Customer Lama' (Existing Customer).\n"
        "  - Memilih Segmen Bisnis: 'Corporate', 'Commercial', atau 'Micro'.\n"
        "  - Menentukan Kantor Cabang LiuGong (misal: JKT01 Jakarta, SBY01 Surabaya, BPN01 Balikpapan, PLB01 Palembang, MDN01 Medan).\n"
        "  - Memverifikasi Tanggal Aplikasi Pembiayaan.\n\n"
        "• Modul 02 - Kantor Registrasi & Domisili (Otomasi Scan OCR NIB 13-Digit):\n"
        "  - Memasukkan Nama Legal Perusahaan (PT/CV).\n"
        "  - Pemanfaatan Fitur AI OCR NIB: Marketing menekan tombol pill '[ 📷 Scan OCR NIB ]' di header field NIB. "
        "Jendela pop-up modal terbuka; Marketing mengunggah berkas sertifikat standar NIB OSS RBA (format PDF atau gambar). "
        "AI OCR langsung membaca berkas dan mengekstrak secara otomatis 5 parameter kunci: 13-Digit Nomor NIB, Nama Perusahaan, Alamat Lengkap, Kode Pos, dan Tanggal Terbit OSS. "
        "Seluruh field langsung terisi otomatis tanpa risiko salah ketik.\n"
        "  - Melengkapi Kota, Provinsi, dan Nomor Telepon/Fax Kantor Pusat.\n\n"
        "• Modul 03 - Profil Usaha & Legalitas Awal:\n"
        "  - Memasukkan Bisnis Utama / KBLI (misal: Kontraktor Penambangan Nikel, Perkebunan Kelapa Sawit, Jasa Konstruksi Sipil).\n"
        "  - Mengisi Pengalaman Usaha (Tahun dan Bulan berjalan).\n"
        "  - Memasukkan Nomor NPWP Badan Usaha (15 atau 16 digit angka).\n"
        "  - Memasukkan Nomor Akta Pendirian dan Tanggal Akta Pendirian awal.\n\n"
        "• Modul 04 - Kontak Utama / Penanggung Jawab (PIC):\n"
        "  - Mengisi Nama Lengkap PIC yang ditunjuk debitur untuk menangani pengadaan alat berat.\n"
        "  - Mengisi Jabatan Formal di Perusahaan (misal: Direktur Operasional, GM Procurement).\n"
        "  - Memasukkan Alamat Domisili Tempat Tinggal PIC, Kota, Provinsi, Kode Pos.\n"
        "  - Memasukkan Nomor Handphone / WhatsApp aktif dan Email kantor resmi PIC.\n\n"
        "• Modul 05 - Segmentasi Industri & Proyek Berjalan:\n"
        "  - Memilih Sektor Industri (Mining Nikel/Batubara, Construction, Forestry, Agriculture, Logistics).\n"
        "  - Mengisi Deskripsi Proyek / Kontrak Kerja yang sedang berjalan (nama bouwheer/pemberi kerja, lokasi tambang/konstruksi, durasi kontrak kerja).\n\n"
        "• Verifikasi Visual Kolom Kanan (Live Excel Mirror 118-Baris):\n"
        "  - Selama Marketing mengetik pada modul kiri, lembar kerja resmi Excel (FAP-Document_APPROVED_v01.2025.xlsx) pada kolom kanan terisi secara real-time.\n"
        "  - Marketing dapat menggunakan tombol navigasi cepat (1. Identitas s/d 5. Proyek) di ribbon Excel untuk memverifikasi tata letak dokumen cetak A4.\n"
        "  - Marketing menekan tombol 'Simpan' untuk mengunci data FAP Stage ke basis data."
    )

    add_image_figure(doc, r"D:\CRM App2\screenshots_presdir\03_workspace_split_50_50.png",
                     "Dual-Pane Split Workspace 50-50: Form Input Accordion Kiri dan Live Digital Twin Excel FAP Kanan (Live System)")

    add_image_figure(doc, r"D:\CRM App2\screenshots_presdir\04_nib_ocr_modal.png",
                     "Fitur Modal Scan OCR Dokumen NIB OSS RBA dengan Ekstraksi 5 Parameter Kunci Otomatis (Live System)")

    # LANGKAH 3
    sub_l3 = doc.add_paragraph()
    sub_l3.paragraph_format.space_before = Pt(8)
    sub_l3.paragraph_format.space_after = Pt(2)
    r = sub_l3.add_run("Langkah 3: Pengisian Tab 2 (Legal Deeds) — Akta Notaris & SK Kemenkumham")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "Setelah data Customer & PIC tersimpan, Marketing beralih ke Tab 2 'Legal':\n"
        "1. Perekaman Akta Pendirian:\n"
        "   - Klik tombol '+ Tambah akta', pilih tipe 'Akta Pendirian'.\n"
        "   - Memasukkan Nomor Akta Notaris, Tanggal Akta, Nama Notaris Pembuat Akta.\n"
        "   - Memasukkan Nomor SK Pengesahan Kemenkumham dan Tanggal SK.\n"
        "2. Perekaman Akta Perubahan Terakhir:\n"
        "   - Menambahkan baris akta perubahan, memilih tipe 'Akta Perubahan'.\n"
        "   - Memasukkan Nomor Akta Perubahan, Tanggal Akta, Nomor dan Tanggal SK Menkumham Perubahan.\n"
        "   - Memilih chip 'Kategori Remarks': Saham, Kepengurusan, Modal, atau Anggaran Dasar.\n"
        "   - Pencatatan Ringkasan Klausul: Menggunakan teks editor TinyMCE dengan penomoran berurutan otomatis (continuous ordered list) untuk mendokumentasikan pasal-pasal penting perubahan akta.\n"
        "3. Verifikasi Status Akta:\n"
        "   - Sistem secara otomatis menempatkan akta yang berlaku efektif pada tabel atas 'Akta Aktif & Rujukan Terkini', "
        "sementara draf atau akta lampau diarsipkan ke tabel bawah 'Riwayat Akta Lama (Historical)'."
    )

    add_image_figure(doc, r"D:\CRM App2\screenshots_presdir\tab_2_legal.png",
                     "Manajemen Akta Legalitas: Tabel Akta Aktif Rujukan Terkini dan Riwayat Akta Historis (Live System)")

    # LANGKAH 4
    sub_l4 = doc.add_paragraph()
    sub_l4.paragraph_format.space_before = Pt(8)
    sub_l4.paragraph_format.space_after = Pt(2)
    r = sub_l4.add_run("Langkah 4: Pengisian Tab 3 (Stakeholders) — Struktur Pengurus & Validasi Saham Mutlak 100%")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "Marketing beralih ke Tab 3 'Stakeholders' untuk mendata jajaran manajemen dan pemegang saham:\n"
        "1. Input Pengurus Perseroan:\n"
        "   - Menambahkan Direktur Utama, Direktur, Komisaris Utama, dan Komisaris.\n"
        "   - Menentukan kewarganegaraan: WNI (wajib memasukkan 16-digit NIK e-KTP) atau WNA (memasukkan Nomor Paspor & izin tinggal E-KITAS untuk direksi ekspatriat mitra LiuGong).\n"
        "2. Input Pemegang Saham & Modal:\n"
        "   - Menambahkan pemilik saham perorangan maupun badan hukum (memasukkan NPWP Perusahaan Pemegang Saham).\n"
        "   - Memasukkan jumlah lembar saham dan nilai nominal per lembar saham.\n"
        "3. Validasi Keras Mutlak 100% (Hard Gate Invariant):\n"
        "   - Sistem menghitung persentase kepemilikan modal secara real-time.\n"
        "   - Total persentase seluruh pemegang saham WAJIB tepat 100.00%. "
        "Jika total saham kurang (misal 99.5%) atau lebih (misal 101.0%), sistem memunculkan indikator peringatan merah dan memblokir tombol pengajuan aplikasi. "
        "Hal ini mengunci kepatuhan hukum kepemilikan perseroan sebelum melangkah ke tahap berikutnya."
    )

    add_image_figure(doc, r"D:\CRM App2\screenshots_presdir\tab_3_stakeholders.png",
                     "Struktur Pemegang Saham & Jajaran Pengurus dengan Validasi Keras Kepemilikan Modal 100% (Live System)")

    # LANGKAH 5
    sub_l5 = doc.add_paragraph()
    sub_l5.paragraph_format.space_before = Pt(8)
    sub_l5.paragraph_format.space_after = Pt(2)
    r = sub_l5.add_run("Langkah 5: Pengisian Tab 4 (Documents) — Verifikasi 13 Berkas Wajib & Unggah Digital")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "Sesuai instruksi Direksi, Tab Dokumen dihadirkan secara penuh pada FAP Stage untuk memastikan berkas fisik terdigitalisasi sejak awal:\n"
        "1. Checklist 13 Dokumen Persyaratan Wajib:\n"
        "   1. KTP Pengurus (Seluruh Direksi & Komisaris)\n"
        "   2. NPWP Badan Usaha & Pengurus\n"
        "   3. NIB (Nomor Induk Berusaha) OSS BKPM\n"
        "   4. Akta Pendirian Perusahaan + SK Pengesahan Kemenkumham\n"
        "   5. Akta Perubahan Terakhir + SK Penerimaan Pemberitahuan Kemenkumham\n"
        "   6. Rekening Koran Operasional Perusahaan (3–6 Bulan Terakhir)\n"
        "   7. Laporan Keuangan Neraca & Laba Rugi (Audited / In-House)\n"
        "   8. SPT Tahunan Pajak Penghasilan Badan\n"
        "   9. Profil Perusahaan (Company Profile & Portofolio Pekerjaan)\n"
        "   10. Salinan Kontrak Kerja / SPK Proyek Berjalan dengan Bouwheer\n"
        "   11. Izin Usaha Sektoral (IUJP/IUP untuk Pertambangan, SBU Konstruksi)\n"
        "   12. Formulir Aplikasi Pembiayaan (FAP) Asli Bertandatangan Basah & Cap Basah Perusahaan\n"
        "   13. Surat Kuasa Persetujuan Pengecekan Data Kredit Pefindo Biro Kredit\n"
        "2. Unggah Berkas Digital: Marketing menekan tombol 'Upload' pada masing-masing dokumen, memilih berkas scan (PDF/JPG) dari komputer.\n"
        "3. Pemantauan Kelengkapan: Status berkas berubah menjadi 'Uploaded/Verified' dan counter kelengkapan di header tab bertambah (misal: 'Documents 13/13')."
    )

    add_image_figure(doc, r"D:\CRM App2\screenshots_presdir\fap_stage_tab_document_view.png",
                     "Pusat Verifikasi Berkas: Checklist 13 Dokumen Wajib dengan Fitur Unggah & Status Kelengkapan (Live System)")

    # LANGKAH 6
    sub_l6 = doc.add_paragraph()
    sub_l6.paragraph_format.space_before = Pt(8)
    sub_l6.paragraph_format.space_after = Pt(2)
    r = sub_l6.add_run("Langkah 6: Final Review & Pengajuan FAP Stage ke Business Support (BS)")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "1. Peninjauan Metrik Header: Marketing memeriksa ringkasan metrik di bagian atas workspace (Nilai Permohonan Pembiayaan, Tenor Angsuran, Kategori Nasabah, dan Status Verifikasi Dokumen 13/13).\n"
        "2. Eksekusi Pengajuan: Marketing menekan tombol 'Submit FAP Stage' (atau 'SUBMIT_TO_BS'). "
        "Sistem melakukan validasi menyeluruh (Nama perusahaan terisi, NIB 13 digit valid, NPWP valid, total saham 100.00%). "
        "Jika seluruh aturan terpenuhi, dialog konfirmasi pengajuan muncul.\n"
        "3. Transisi Tahap Otomatis: Status aplikasi secara resmi berpindah dari 'FAP Stage (Draft)' menjadi 'SUBMITTED_TO_BS'. "
        "Aplikasi diteruskan secara digital ke meja kerja Business Support (BS) untuk verifikasi awal, dan tahap 'Pre-analysis' diaktifkan. "
        "Seluruh riwayat aksi tercatat permanen pada tab Activity Log (Audit Trail abadi)."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # BAB 3: GAMBARAN UMUM FASE MENDATANG (ROADMAP 2027)
    # -------------------------------------------------------------
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    r = h3.add_run("BAB 3: GAMBARAN UMUM FASE MENDATANG (ROADMAP 2027)")
    r.font.name = "Calibri"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Fase lanjutan pada tahun 2027 akan mengembangkan kapabilitas analitis mendalam, "
        "decision engine otomatis, serta penuntasan proposal komite kredit:"
    )

    # 3.1 Q1 2027
    sub31 = doc.add_paragraph()
    sub31.paragraph_format.space_before = Pt(6)
    sub31.paragraph_format.space_after = Pt(2)
    r = sub31.add_run("3.1. Q1 2027 — Pre-Analysis & Integrasi Pefindo OCR (OCR NPWP / SPT Tahunan)")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "• Early Risk Triaging: Skrining kelayakan awal debitur secara instan (kategori Go / Review / No-Go) sebelum analis menghabiskan waktu ke lapangan.\n"
        "• Integrasi Pefindo Biro Kredit: Menembakkan identitas nasabah ke Pefindo untuk menarik Credit Score (skala 250–900), riwayat pembayaran kredit, serta total fasilitas utang terbuka secara real-time.\n"
        "• Fitur Cerdas OCR NPWP & SPT Tahunan: Pemindai OCR mutakhir yang membaca lampiran SPT Pajak Penghasilan Badan (Formulir 1771) dan secara langsung memetakan angka Penjualan, HPP, Beban Operasional, Aset Lancar, dan Modal ke tabel Financial Statement di sistem tanpa pengetikan manual."
    )

    # 3.2 Q2 2027
    sub32 = doc.add_paragraph()
    sub32.paragraph_format.space_before = Pt(6)
    sub32.paragraph_format.space_after = Pt(2)
    r = sub32.add_run("3.2. Q2 2027 — ICLF & Credit Decision Engine (CDE Financial)")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "• Initial Credit Limit Facility (ICLF): Penetapan batas pagu pembiayaan maksimum per debitur tunggal maupun grup usaha berdasarkan kepatuhan Batas Maksimum Pemberian Pembiayaan (BMPP OJK) dan rasio permodalan.\n"
        "• Credit Decision Engine (CDE) Phase 1: Engine penilaian kredit otomatis yang mengevaluasi rasio keuangan (Current Ratio, Debt-to-Equity, Net Profit Margin, Interest Coverage Ratio) dan menghasilkan skor kelayakan finansial secara terstandardisasi."
    )

    # 3.3 Q3 2027
    sub33 = doc.add_paragraph()
    sub33.paragraph_format.space_before = Pt(6)
    sub33.paragraph_format.space_after = Pt(2)
    r = sub33.add_run("3.3. Q3 2027 — Credit Decision Engine (CDE Bank Statement) Phase 2")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "• Parser Rekening Koran Cerdas: Mengekstrak mutasi buku bank debitur (Bank Mandiri, BCA, BRI, BNI) dari berkas PDF mutasi 3–6 bulan terakhir.\n"
        "• Analisis Likuiditas Riil: Menghitung saldo rata-rata bulanan, perputaran kredit harian, frekuensi saldo minus, serta mendeteksi transaksi anomali atau cek kosong untuk memastikan debitur memiliki kemampuan membayar angsuran riil."
    )

    # 3.4 Q4 2027
    sub34 = doc.add_paragraph()
    sub34.paragraph_format.space_before = Pt(6)
    sub34.paragraph_format.space_after = Pt(2)
    r = sub34.add_run("3.4. Q4 2027 — Proposal Pembiayaan Final & Putusan Komite Kredit")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "• Rincian Unit Alat Berat LiuGong: Inventaris unit (Excavator 922E, Wheel Loader 856H, Dump Truck), harga OTR, subsidi dealer, dan uang muka (DP 20%–30%).\n"
        "• Mesin Aktuaria Bunga Efektif: Simulasi angsuran anuitas menurun resmi (reducing balance annuity) presisi tinggi.\n"
        "• Persetujuan Komite Kredit Digital: Pelaksanaan sidang komite kredit digital berjenjang (Tier 1 Cabang s/d Tier 3 Direksi/Presdir) dengan hak Approve, Reject, atau Return, disertai pencatatan audit trail permanen."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # BAB 4: LEMBAR PENGESAHAN EKSEKUTIF
    # -------------------------------------------------------------
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(4)
    r = h4.add_run("BAB 4: LEMBAR PENGESAHAN EKSEKUTIF")
    r.font.name = "Calibri"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Proposal Spesifikasi Bisnis dan Roadmap Sistem LOS-CRM ini disusun sebagai pedoman resmi pengembangan sistem "
        "dan diajukan untuk memperoleh persetujuan prinsip dari Bapak Presiden Direktur:"
    )

    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    set_table_borders(sign_table, color="D3D3D3")
    col_w_sign = [Inches(3.2), Inches(3.2)]
    
    sign_headers = ["Diajukan Oleh:", "Disetujui Oleh:"]
    for j, htext in enumerate(sign_headers):
        c = sign_table.rows[0].cells[j]
        c.width = col_w_sign[j]
        set_cell_background(c, "002060")
        set_cell_margins(c, 50, 50, 80, 80)
        p_hdr = c.paragraphs[0]
        r_hdr = p_hdr.add_run(htext)
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(9.5)
        r_hdr.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)

    sign_contents = [
        ("Latom\nLead System Owner & AI CTO\nPT LiuGong Finance Indonesia\n\nTanggal: 19 Oktober 2026"),
        ("Presiden Direktur\nChief Executive Officer\nPT LiuGong Finance Indonesia\n\nTanggal: .............................")
    ]
    for j, text in enumerate(sign_contents):
        c = sign_table.rows[1].cells[j]
        c.width = col_w_sign[j]
        set_cell_margins(c, 100, 100, 80, 80)
        set_cell_background(c, "FFFFFF")
        p_c = c.paragraphs[0]
        p_c.paragraph_format.line_spacing = 1.2
        r_c = p_c.add_run(text)
        r_c.font.name = "Calibri"
        r_c.font.size = Pt(9.5)
        r_c.font.color.rgb = RGBColor(0, 32, 96)

    # Simpan dokumen
    out_path = r"D:\CRM App2\CRM App2\LIUGONG_CRM_BUSINESS_SPECIFICATION_PRESDIR.docx"
    doc.save(out_path)
    print(f"Step-by-step SOP Document created at: {out_path} (size: {os.path.getsize(out_path)} bytes)")

if __name__ == "__main__":
    build_word_document()
